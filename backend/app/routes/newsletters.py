"""Newsletters: sending account, contacts and lists, campaigns, and public tracking links."""
from __future__ import annotations

import asyncio
import csv
import io
import re
from datetime import datetime, timezone

from fastapi import APIRouter, BackgroundTasks, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, Response
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..auth import RequireOwner, RequireUser, RequireWriter
from ..db import Brand, Contact, ContactList, Delivery, Newsletter, current_org, get_db, session_scope, utcnow
from ..services import generator, mailer, newsletter, newsletter_types, pipeline

router = APIRouter(prefix="/api", tags=["newsletters"], dependencies=[RequireUser, RequireWriter])
public = APIRouter(prefix="/api/t", tags=["tracking"])
EMAIL = re.compile(r"^[^@\s,;<>\"']+@[^@\s,;<>\"']+\.[a-zA-Z]{2,}$")


# ------------------------------------------------------------------ sending account
@router.get("/email/settings")
def email_settings(db: Session = Depends(get_db)):
    return mailer.public(db)


@router.put("/email/settings", dependencies=[RequireOwner])
def save_email_settings(body: dict, db: Session = Depends(get_db)):
    if body.get("from_email") and not EMAIL.match(str(body["from_email"]).strip()):
        raise HTTPException(400, "بريد إلكتروني غير صالح")
    return mailer.save(db, body)


class TestMail(BaseModel):
    to: str


@router.post("/email/test")
async def test_email(body: TestMail, db: Session = Depends(get_db)):
    cfg = mailer.load(db)
    if not EMAIL.match(body.to.strip()):
        raise HTTPException(400, "بريد إلكتروني غير صالح")
    html = "<p style='font-family:Arial'>✅ Newsletter sending works.<br>✅ الإرسال يعمل بنجاح.</p>"
    try:
        await asyncio.to_thread(mailer.send, cfg, body.to.strip(), "", "Test ✓", html, "Newsletter sending works.")
    except mailer.MailError as exc:
        raise HTTPException(400, str(exc)) from exc
    return {"ok": True}


# ------------------------------------------------------------------ lists
class ListIn(BaseModel):
    name: str = Field(min_length=1, max_length=160)
    description: str = ""


def _list_out(db: Session, lst: ContactList) -> dict:
    members = sum(1 for c in db.scalars(select(Contact)) if lst.id in (c.list_ids or []))
    return {**lst.to_dict(), "members": members}


@router.get("/contacts/lists")
def lists(db: Session = Depends(get_db)):
    contacts = list(db.scalars(select(Contact)))
    out = []
    for lst in db.scalars(select(ContactList).order_by(ContactList.id)):
        members = [c for c in contacts if lst.id in (c.list_ids or [])]
        out.append({**lst.to_dict(), "members": len(members),
                    "subscribed": sum(1 for c in members if c.status == "subscribed")})
    return out


@router.post("/contacts/lists")
def create_list(body: ListIn, db: Session = Depends(get_db)):
    lst = ContactList(name=body.name.strip(), description=body.description)
    db.add(lst)
    db.flush()
    return _list_out(db, lst)


@router.patch("/contacts/lists/{lid}")
def update_list(lid: int, body: ListIn, db: Session = Depends(get_db)):
    lst = db.get(ContactList, lid)
    if lst is None:
        raise HTTPException(404, "القائمة غير موجودة")
    lst.name, lst.description = body.name.strip(), body.description
    return _list_out(db, lst)


@router.delete("/contacts/lists/{lid}")
def delete_list(lid: int, delete_contacts: bool = False, db: Session = Depends(get_db)):
    lst = db.get(ContactList, lid)
    if lst is None:
        raise HTTPException(404, "القائمة غير موجودة")
    for c in db.scalars(select(Contact)):
        if lid in (c.list_ids or []):
            rest = [x for x in c.list_ids if x != lid]
            if delete_contacts and not rest:
                db.delete(c)
            else:
                c.list_ids = rest
    db.delete(lst)
    return {"ok": True}


# ------------------------------------------------------------------ contacts
@router.get("/contacts")
def contacts(list_id: int | None = None, q: str = "", status: str = "", limit: int = 300, offset: int = 0,
             db: Session = Depends(get_db)):
    rows = list(db.scalars(select(Contact).order_by(Contact.id.desc())))
    if list_id:
        rows = [c for c in rows if list_id in (c.list_ids or [])]
    if status:
        rows = [c for c in rows if c.status == status]
    if q.strip():
        needle = q.strip().lower()
        rows = [c for c in rows if needle in c.email.lower() or needle in (c.name or "").lower()]
    return {"total": len(rows), "items": [c.to_dict() for c in rows[offset:offset + min(limit, 1000)]]}


class ContactIn(BaseModel):
    email: str
    name: str = ""
    list_ids: list[int] = []


def _upsert(db: Session, email: str, name: str, list_ids: list[int], source: str, fields: dict | None = None) -> str:
    email = email.strip().lower()
    if not EMAIL.match(email):
        return "invalid"
    c = db.scalar(select(Contact).where(Contact.email == email))
    if c is None:
        db.add(Contact(email=email, name=name.strip()[:200], list_ids=list(dict.fromkeys(list_ids)), source=source,
                       fields=fields or {}))
        return "added"
    c.list_ids = list(dict.fromkeys([*(c.list_ids or []), *list_ids]))
    if name and not c.name:
        c.name = name.strip()[:200]
    if fields:
        c.fields = {**(c.fields or {}), **fields}
    return "updated"


@router.post("/contacts")
def add_contact(body: ContactIn, db: Session = Depends(get_db)):
    result = _upsert(db, body.email, body.name, body.list_ids, "manual")
    if result == "invalid":
        raise HTTPException(400, "بريد إلكتروني غير صالح")
    db.flush()
    return {"result": result}


class ContactPatch(BaseModel):
    name: str | None = None
    status: str | None = Field(None, pattern="^(subscribed|unsubscribed)$")
    list_ids: list[int] | None = None


@router.patch("/contacts/{cid}")
def update_contact(cid: int, body: ContactPatch, db: Session = Depends(get_db)):
    c = db.get(Contact, cid)
    if c is None:
        raise HTTPException(404, "جهة الاتصال غير موجودة")
    if body.name is not None:
        c.name = body.name.strip()
    if body.status:
        c.status = body.status
        c.unsubscribed_at = utcnow() if body.status == "unsubscribed" else None
    if body.list_ids is not None:
        c.list_ids = body.list_ids
    return c.to_dict()


class ContactBulk(BaseModel):
    ids: list[int] = Field(min_length=1)
    action: str = Field(pattern="^(delete|unsubscribe|resubscribe|add_to_list|remove_from_list)$")
    list_id: int | None = None


@router.post("/contacts/bulk")
def contacts_bulk(body: ContactBulk, db: Session = Depends(get_db)):
    rows = list(db.scalars(select(Contact).where(Contact.id.in_(body.ids))))
    for c in rows:
        if body.action == "delete":
            db.delete(c)
        elif body.action == "unsubscribe":
            c.status, c.unsubscribed_at = "unsubscribed", utcnow()
        elif body.action == "resubscribe":
            c.status, c.unsubscribed_at = "subscribed", None
        elif body.action == "add_to_list" and body.list_id:
            c.list_ids = list(dict.fromkeys([*(c.list_ids or []), body.list_id]))
        elif body.action == "remove_from_list" and body.list_id:
            c.list_ids = [x for x in (c.list_ids or []) if x != body.list_id]
    return {"done": len(rows)}


def _rows_from_file(name: str, data: bytes) -> list[list[str]]:
    name = name.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        return [[("" if v is None else str(v)).strip() for v in row] for row in wb.active.iter_rows(values_only=True)]
    text = None
    for enc in ("utf-8-sig", "utf-16", "cp1256", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    text = text or ""
    if name.endswith(".vcf"):
        rows, current = [], ["", ""]
        for line in text.splitlines():
            if line.upper().startswith("FN"):
                current[1] = line.split(":", 1)[-1].strip()
            elif "EMAIL" in line.upper() and ":" in line:
                current[0] = line.split(":", 1)[-1].strip()
            elif line.upper().startswith("END:VCARD"):
                rows.append(current)
                current = ["", ""]
        return [["email", "name"], *rows]
    sample = text[:4000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]


def _import_rows(db: Session, rows: list[list[str]], list_ids: list[int]) -> dict:
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return {"added": 0, "updated": 0, "invalid": 0}
    header = [c.lower() for c in rows[0]]
    has_header = not any(EMAIL.match(c) for c in rows[0])
    email_col = next((i for i, h in enumerate(header) if re.search(r"e-?mail|بريد|ایمیل|courriel|emel", h)), None)
    name_col = next((i for i, h in enumerate(header) if re.search(r"^(full ?)?name|الاسم|اسم|nom|nama", h)), None)
    first_col = next((i for i, h in enumerate(header) if re.search(r"first|prénom|الأول", h)), None)
    last_col = next((i for i, h in enumerate(header) if re.search(r"last|nom de famille|العائلة", h)), None)
    body = rows[1:] if has_header else rows
    if email_col is None:          # find the column that holds e-mails
        counts = {}
        for r in body[:50]:
            for i, c in enumerate(r):
                if EMAIL.match(c.strip()):
                    counts[i] = counts.get(i, 0) + 1
        email_col = max(counts, key=counts.get) if counts else 0
    result = {"added": 0, "updated": 0, "invalid": 0}
    for r in body:
        email = r[email_col] if email_col < len(r) else ""
        if name_col is not None and name_col < len(r):
            name = r[name_col]
        else:
            name = " ".join(r[i] for i in (first_col, last_col) if i is not None and i < len(r) and r[i])
        extra = {rows[0][i]: v for i, v in enumerate(r) if has_header and i not in (email_col, name_col) and v
                 and i < len(rows[0])} if has_header else {}
        res = _upsert(db, email, name, list_ids, "import", dict(list(extra.items())[:10]))
        result[res] += 1
    return result


def _target_lists(db: Session, list_id: int | None, new_list: str) -> list[int]:
    if new_list.strip():
        lst = ContactList(name=new_list.strip()[:160])
        db.add(lst)
        db.flush()
        return [lst.id]
    return [list_id] if list_id else []


@router.post("/contacts/import")
async def import_contacts(file: UploadFile = File(...), list_id: int | None = Form(None),
                          new_list: str = Form(""), db: Session = Depends(get_db)):
    """Excel (.xlsx), CSV/TSV, TXT (one address per line) or vCard (.vcf)."""
    data = await file.read()
    if len(data) > 15 * 1024 * 1024:
        raise HTTPException(400, "الحد الأقصى 15 ميغابايت")
    try:
        rows = await asyncio.to_thread(_rows_from_file, file.filename or "", data)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(400, "تعذّر قراءة الملف — استخدم Excel أو CSV أو TXT") from exc
    return _import_rows(db, rows, _target_lists(db, list_id, new_list))


class PasteIn(BaseModel):
    text: str
    list_id: int | None = None
    new_list: str = ""


@router.post("/contacts/import-text")
def import_text(body: PasteIn, db: Session = Depends(get_db)):
    """Pasted addresses: one per line, or "Name <email>", or separated by commas."""
    rows = [["email", "name"]]
    for part in re.split(r"[\n,;]+", body.text):
        m = re.search(r"([^<\s]+@[^>\s]+)", part)
        if m:
            name = part.replace(m.group(1), "").strip(" <>\"'\t")
            rows.append([m.group(1), name])
    return _import_rows(db, rows, _target_lists(db, body.list_id, body.new_list))


@router.get("/contacts/export.csv")
def export_contacts(list_id: int | None = None, db: Session = Depends(get_db)):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["email", "name", "status", "lists", "created_at"])
    names = {lst.id: lst.name for lst in db.scalars(select(ContactList))}
    for c in db.scalars(select(Contact).order_by(Contact.id)):
        if list_id and list_id not in (c.list_ids or []):
            continue
        w.writerow([c.email, c.name, c.status, "|".join(names.get(i, "") for i in c.list_ids or []),
                    c.created_at.isoformat() if c.created_at else ""])
    return Response("﻿" + buf.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": "attachment; filename=contacts.csv"})


# ------------------------------------------------------------------ newsletters
def _family(db: Session) -> str:
    return pipeline.family_of(db, current_org.get())


def _nl(db: Session, nid: int) -> Newsletter:
    nl = db.get(Newsletter, nid)
    if nl is None:
        raise HTTPException(404, "النشرة غير موجودة")
    return nl


def _out(db: Session, nl: Newsletter, full: bool = False) -> dict:
    out = {**nl.to_dict(), "stats": newsletter.stats(db, nl),
           "audience": len(newsletter.recipients(db, nl.list_ids))}
    if not full:
        out.pop("content", None)
    return out


@router.get("/newsletters/fonts")
def newsletter_fonts(language: str = "ar"):
    return newsletter.font_options(language if language in ("ar", "en", "ms", "fr") else "ar")


@router.get("/newsletters/designs")
def designs(db: Session = Depends(get_db)):
    return newsletter.designs_for(_family(db))


@router.get("/newsletters/types")
def types():
    """The kinds of newsletter (curated, educational, promotional, event…)."""
    return newsletter_types.options()


@router.get("/newsletters/articles")
def recent_articles(db: Session = Depends(get_db)):
    """Recently collected articles from the company's sources — material for curated / news emails."""
    from ..db import Article, Source
    names = {s.id: s.name for s in db.scalars(select(Source))}
    rows = db.scalars(select(Article).where(Article.status != "error").order_by(Article.fetched_at.desc()).limit(60))
    return [{"id": a.id, "title": a.title, "url": a.url, "image_url": a.image_url,
             "source": names.get(a.source_id, ""), "published_at": a.to_dict()["published_at"]} for a in rows]


@router.get("/newsletters")
def list_newsletters(db: Session = Depends(get_db)):
    return [_out(db, n) for n in db.scalars(select(Newsletter).order_by(Newsletter.id.desc()))]


class EventIn(BaseModel):
    date: str = ""
    time: str = ""
    place: str = ""
    url: str = ""


class PollIn(BaseModel):
    question: str = ""
    options: list[str] = []


class NewsletterIn(BaseModel):
    subject: str = ""
    kind: str = "hybrid"
    design: str | None = None
    article_ids: list[int] = []
    event: EventIn | None = None
    poll: PollIn | None = None
    language: str | None = Field(None, pattern="^(ar|en|ms|fr)$")
    font: str = ""
    list_ids: list[int] = []
    topic: str = ""
    draft_ids: list[int] = []
    product_ids: list[int] = []
    sections: int = Field(0, ge=0, le=8)          # 0 = what the newsletter type usually has


@router.post("/newsletters")
def create_newsletter(body: NewsletterIn, background: BackgroundTasks, db: Session = Depends(get_db)):
    family = _family(db)
    kind = body.kind if body.kind in newsletter_types.TYPES else "hybrid"
    field_designs = newsletter.FIELD_LAYOUTS.get(family, ["classic"])
    suggested = newsletter_types.get(kind)["design"]
    design = body.design if body.design in newsletter.LAYOUTS else (
        suggested if suggested in field_designs else field_designs[0])
    lang = body.language or newsletter.org_language(db, current_org.get())
    extra = {}
    if body.event and any(body.event.model_dump().values()):
        extra["event"] = {k: v.strip() for k, v in body.event.model_dump().items() if v.strip()}
    if body.poll and body.poll.question.strip() and [o for o in body.poll.options if o.strip()]:
        extra["poll"] = {"question": body.poll.question.strip(),
                         "options": [o.strip() for o in body.poll.options if o.strip()][:6]}
    font = body.font if body.font in newsletter.EMAIL_FONTS else ""
    nl = Newsletter(subject=body.subject.strip() or body.topic.strip(), design=design, language=lang, kind=kind,
                    font=font, list_ids=body.list_ids, status="generating", content={})
    db.add(nl)
    db.flush()
    nid = nl.id
    db.commit()

    async def run() -> None:
        try:
            await newsletter.generate(nid, body.topic, body.draft_ids, body.product_ids, body.sections,
                                      body.article_ids, extra)
        except Exception as exc:  # noqa: BLE001
            with session_scope() as s:
                n = s.get(Newsletter, nid)
                n.status, n.error = "draft", f"فشل توليد النص: {exc}"[:1000]
    background.add_task(run)
    return {"id": nid}


@router.get("/newsletters/{nid}")
def get_newsletter(nid: int, db: Session = Depends(get_db)):
    return _out(db, _nl(db, nid), full=True)


class NewsletterPatch(BaseModel):
    kind: str | None = None
    subject: str | None = None
    preheader: str | None = None
    design: str | None = None
    content: dict | None = None
    list_ids: list[int] | None = None
    language: str | None = Field(None, pattern="^(ar|en|ms|fr)$")
    font: str | None = None


@router.patch("/newsletters/{nid}")
def update_newsletter(nid: int, body: NewsletterPatch, db: Session = Depends(get_db)):
    nl = _nl(db, nid)
    if nl.status in ("sending", "sent"):
        raise HTTPException(409, "لا يمكن تعديل نشرة أُرسلت")
    data = body.model_dump(exclude_unset=True)
    if data.get("design") and data["design"] not in newsletter.LAYOUTS:
        data.pop("design")
    if data.get("kind") and data["kind"] not in newsletter_types.TYPES:
        data.pop("kind")
    if "font" in data and data["font"] not in newsletter.EMAIL_FONTS:
        data["font"] = ""
    for k, v in data.items():
        if v is not None:
            setattr(nl, k, v)
    return _out(db, nl, full=True)


class TranslateIn(BaseModel):
    language: str = Field(..., pattern="^(ar|en|ms|fr)$")


@router.post("/newsletters/{nid}/translate")
def translate_newsletter(nid: int, body: TranslateIn, background: BackgroundTasks, db: Session = Depends(get_db)):
    nl = _nl(db, nid)
    if nl.status in ("sending", "sent", "generating"):
        raise HTTPException(409, "لا يمكن تعديل هذه النشرة الآن")
    if not generator.ai_available():
        raise HTTPException(400, "الذكاء الاصطناعي غير مفعّل — أضف مفتاحًا من الإعدادات")
    nl.status, nl.error = "generating", None
    db.commit()

    async def run() -> None:
        try:
            await newsletter.translate(nid, body.language)
        except Exception as exc:  # noqa: BLE001
            with session_scope() as s:
                n = s.get(Newsletter, nid)
                n.status, n.error = "draft", f"فشلت الترجمة: {exc}"[:1000]
    background.add_task(run)
    return {"ok": True}


@router.delete("/newsletters/{nid}")
def delete_newsletter(nid: int, db: Session = Depends(get_db)):
    nl = _nl(db, nid)
    if nl.status == "sending":
        raise HTTPException(409, "النشرة قيد الإرسال الآن")
    for d in db.scalars(select(Delivery).where(Delivery.newsletter_id == nid)):
        db.delete(d)
    db.delete(nl)
    return {"ok": True}


@router.post("/newsletters/{nid}/duplicate")
def duplicate(nid: int, db: Session = Depends(get_db)):
    src = _nl(db, nid)
    nl = Newsletter(subject=src.subject, preheader=src.preheader, design=src.design, language=src.language,
                    kind=src.kind, font=src.font or "",
                    content=dict(src.content or {}), list_ids=list(src.list_ids or []), status="draft")
    db.add(nl)
    db.flush()
    return {"id": nl.id}


@router.get("/newsletters/{nid}/preview")
def preview(nid: int, design: str | None = None, db: Session = Depends(get_db)):
    nl = _nl(db, nid)
    brand = db.scalar(select(Brand).order_by(Brand.is_default.desc(), Brand.id))
    if design in newsletter.LAYOUTS:
        nl.design = design
        db.expunge(nl)
    html, _ = newsletter.render(nl, brand, _family(db))
    return {"html": html}


@router.post("/newsletters/{nid}/preview")
def live_preview(nid: int, body: NewsletterPatch, db: Session = Depends(get_db)):
    """Preview unsaved edits: the changes are applied to a detached copy and never stored."""
    nl = _nl(db, nid)
    brand = db.scalar(select(Brand).order_by(Brand.is_default.desc(), Brand.id))
    db.expunge(nl)
    for k, v in body.model_dump(exclude_unset=True).items():
        if v is None or (k == "design" and v not in newsletter.LAYOUTS):
            continue
        if k == "font" and v not in newsletter.EMAIL_FONTS:
            v = ""
        setattr(nl, k, v)
    html, _ = newsletter.render(nl, brand, _family(db))
    return {"html": html}


@router.post("/newsletters/{nid}/test")
async def send_test(nid: int, body: TestMail, db: Session = Depends(get_db)):
    nl = _nl(db, nid)
    brand = db.scalar(select(Brand).order_by(Brand.is_default.desc(), Brand.id))
    if not EMAIL.match(body.to.strip()):
        raise HTTPException(400, "بريد إلكتروني غير صالح")
    html, text = newsletter.render(nl, brand, _family(db))
    try:
        await asyncio.to_thread(mailer.send, mailer.load(db), body.to.strip(), "", f"[TEST] {nl.subject}", html, text)
    except mailer.MailError as exc:
        raise HTTPException(400, str(exc)) from exc
    return {"ok": True}


class SendIn(BaseModel):
    when: datetime | None = None          # empty = now


@router.post("/newsletters/{nid}/send")
def send(nid: int, body: SendIn, background: BackgroundTasks, db: Session = Depends(get_db)):
    nl = _nl(db, nid)
    if nl.status in ("sending", "sent"):
        raise HTTPException(409, "النشرة أُرسلت أو قيد الإرسال")
    if not (nl.subject or "").strip() or not (nl.content or {}).get("headline") and not (nl.content or {}).get("sections"):
        raise HTTPException(400, "أكمل عنوان النشرة ومحتواها أولًا")
    if not mailer.ready(mailer.load(db)):
        raise HTTPException(400, "حساب الإرسال غير مضبوط — أكمل إعدادات البريد أولًا")
    if not newsletter.recipients(db, nl.list_ids):
        raise HTTPException(400, "لا يوجد مشتركون في القوائم المختارة")
    when = body.when
    if when and when.tzinfo is None:
        when = when.replace(tzinfo=timezone.utc)
    if when and when > datetime.now(timezone.utc):
        nl.status, nl.scheduled_at = "scheduled", when
        return _out(db, nl)
    nl.status, nl.scheduled_at = "sending", None
    db.commit()
    background.add_task(asyncio.to_thread, newsletter.send_newsletter, nid, current_org.get())
    return _out(db, nl)


@router.post("/newsletters/{nid}/cancel")
def cancel(nid: int, db: Session = Depends(get_db)):
    nl = _nl(db, nid)
    if nl.status != "scheduled":
        raise HTTPException(409, "النشرة غير مجدولة")
    nl.status, nl.scheduled_at = "draft", None
    return _out(db, nl)


@router.get("/newsletters/{nid}/recipients")
def deliveries(nid: int, filter: str = "", db: Session = Depends(get_db)):  # noqa: A002
    _nl(db, nid)
    rows = list(db.scalars(select(Delivery).where(Delivery.newsletter_id == nid).order_by(Delivery.id)))
    if filter == "opened":
        rows = [d for d in rows if d.opened_at]
    elif filter == "clicked":
        rows = [d for d in rows if d.clicked_at]
    elif filter == "unopened":
        rows = [d for d in rows if d.status == "sent" and not d.opened_at]
    elif filter == "failed":
        rows = [d for d in rows if d.status == "failed"]
    elif filter == "unsubscribed":
        rows = [d for d in rows if d.unsubscribed_at]
    return [{k: v for k, v in d.to_dict().items() if k != "token"} for d in rows[:2000]]


# ------------------------------------------------------------------ public tracking (no sign-in)
PIXEL = bytes.fromhex("47494638396101000100800000000000ffffff21f90401000000002c00000000010001000002024401003b")


def _delivery(db: Session, token: str) -> Delivery | None:
    return db.scalar(select(Delivery).where(Delivery.token == token).execution_options(all_orgs=True))


@public.get("/o/{token}.gif")
def opened(token: str):
    with session_scope() as db:
        d = _delivery(db, token)
        if d:
            d.open_count = (d.open_count or 0) + 1
            d.opened_at = d.opened_at or utcnow()
    return Response(PIXEL, media_type="image/gif", headers={"Cache-Control": "no-store, max-age=0"})


@public.get("/c/{token}")
def clicked(token: str, u: str = ""):
    if not u.startswith(("http://", "https://")):
        raise HTTPException(400, "رابط غير صالح")
    with session_scope() as db:
        d = _delivery(db, token)
        if d:
            d.click_count = (d.click_count or 0) + 1
            d.clicked_at = d.clicked_at or utcnow()
            d.opened_at = d.opened_at or utcnow()          # a click means it was opened
    return RedirectResponse(u, status_code=302)


THANKS_PAGE = {
    "ar": ("شكرًا لمشاركتك!", "سُجّل رأيك.", "rtl"), "en": ("Thanks for your answer!", "Your vote was recorded.", "ltr"),
    "ms": ("Terima kasih atas jawapan anda!", "Undian anda telah direkodkan.", "ltr"),
    "fr": ("Merci pour votre réponse !", "Votre vote a été enregistré.", "ltr"),
}


@public.get("/p/{token}", response_class=HTMLResponse)
def poll_answer(token: str, a: int = 0):
    lang = "ar"
    with session_scope() as db:
        d = _delivery(db, token)
        if d:
            nl = db.get(Newsletter, d.newsletter_id)
            lang = nl.language if nl and nl.language in THANKS_PAGE else "ar"
            options = ((nl.content or {}).get("poll") or {}).get("options") or [] if nl else []
            if 0 <= a < len(options):
                d.answer = options[a][:200]
                d.clicked_at = d.clicked_at or utcnow()
                d.opened_at = d.opened_at or utcnow()
    title, text, dir_ = THANKS_PAGE[lang]
    return HTMLResponse(f'<!doctype html><html dir="{dir_}"><head><meta charset="utf-8">'
                        f'<meta name="viewport" content="width=device-width,initial-scale=1"><title>{title}</title></head>'
                        f'<body style="font-family:Tahoma,Arial,sans-serif;background:#f4f4f4;display:flex;'
                        f'align-items:center;justify-content:center;height:100vh;margin:0"><div style="background:#fff;'
                        f'padding:40px;border-radius:16px;text-align:center;max-width:420px"><h2>{title}</h2>'
                        f'<p style="color:#555">{text}</p></div></body></html>')


UNSUB_PAGE = {
    "ar": ("تم إلغاء اشتراكك", "لن تصلك رسائل أخرى من هذه القائمة.", "rtl"),
    "en": ("You have been unsubscribed", "You will not receive more emails from this list.", "ltr"),
    "ms": ("Anda telah dinyahlanggan", "Anda tidak akan menerima e-mel lagi daripada senarai ini.", "ltr"),
    "fr": ("Vous êtes désabonné", "Vous ne recevrez plus d’e-mails de cette liste.", "ltr"),
}


def _unsubscribe(token: str) -> str:
    lang = "ar"
    with session_scope() as db:
        d = _delivery(db, token)
        if d:
            d.unsubscribed_at = d.unsubscribed_at or utcnow()
            nl = db.get(Newsletter, d.newsletter_id)
            lang = nl.language if nl and nl.language in UNSUB_PAGE else "ar"
            c = db.scalar(select(Contact).where(Contact.id == d.contact_id).execution_options(all_orgs=True))
            if c:
                c.status, c.unsubscribed_at = "unsubscribed", utcnow()
    return lang


@public.get("/u/{token}", response_class=HTMLResponse)
def unsubscribe_page(token: str):
    title, text, dir_ = UNSUB_PAGE[_unsubscribe(token)]
    return HTMLResponse(f'<!doctype html><html dir="{dir_}"><head><meta charset="utf-8">'
                        f'<meta name="viewport" content="width=device-width,initial-scale=1"><title>{title}</title></head>'
                        f'<body style="font-family:Tahoma,Arial,sans-serif;background:#f4f4f4;display:flex;'
                        f'align-items:center;justify-content:center;height:100vh;margin:0"><div style="background:#fff;'
                        f'padding:40px;border-radius:16px;text-align:center;max-width:420px"><h2>{title}</h2>'
                        f'<p style="color:#555">{text}</p></div></body></html>')


@public.post("/u/{token}")
def unsubscribe_one_click(token: str):
    _unsubscribe(token)
    return {"ok": True}
